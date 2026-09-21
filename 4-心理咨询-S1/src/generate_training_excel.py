#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
训练记录Excel导出模块
功能：将训练记录导出为专业的Excel报告（两个Sheet）
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from pathlib import Path
import json
from datetime import datetime

def generate_training_excel(session_data: dict) -> str:
    """
    生成训练记录Excel文件

    参数:
        session_data: 训练会话数据（完整JSON）

    返回:
        生成的Excel文件路径
    """
    wb = Workbook()

    # 删除默认的Sheet
    wb.remove(wb.active)

    # Sheet1: 来访者背景资料
    ws1 = wb.create_sheet("来访者背景资料", 0)
    _write_persona_sheet(ws1, session_data)

    # Sheet2: 咨询过程详细分析
    ws2 = wb.create_sheet("咨询过程分析", 1)
    _write_dialogue_sheet(ws2, session_data)

    # 保存文件
    output_dir = Path(__file__).parent.parent / 'output' / '对话模拟训练' / 'exports'
    output_dir.mkdir(parents=True, exist_ok=True)

    session_id = session_data['session_id']
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'训练记录_{session_id}_{timestamp}.xlsx'
    file_path = output_dir / filename

    wb.save(str(file_path))

    return str(file_path)

def _write_persona_sheet(ws, session_data):
    """填写Sheet1：来访者背景资料"""
    persona = session_data['visitor_persona']
    approach = session_data['approach']
    created_at = session_data['created_at']

    # 样式定义
    title_font = Font(name='微软雅黑', size=16, bold=True, color='FFFFFF')
    title_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(name='微软雅黑', size=11, bold=True)
    header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    content_font = Font(name='微软雅黑', size=10)

    # 设置列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 60

    row = 1

    # 标题
    ws.merge_cells(f'A{row}:B{row}')
    cell = ws[f'A{row}']
    cell.value = '对话模拟训练 - 来访者背景资料'
    cell.font = title_font
    cell.fill = title_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 30
    row += 2

    # 基本信息
    _add_row(ws, row, '训练日期', created_at[:19].replace('T', ' '), header_font, header_fill, content_font)
    row += 1
    _add_row(ws, row, '督导流派', approach, header_font, header_fill, content_font)
    row += 2

    # 来访者信息
    _add_row(ws, row, '姓名', persona.get('name', ''), header_font, header_fill, content_font)
    row += 1
    _add_row(ws, row, '年龄/性别', f"{persona.get('age', '')}岁 / {persona.get('gender', '')}", header_font, header_fill, content_font)
    row += 1
    _add_row(ws, row, '情绪状态', persona.get('emotional_state', ''), header_font, header_fill, content_font)
    row += 2

    # 背景概要
    ws.merge_cells(f'A{row}:B{row}')
    cell = ws[f'A{row}']
    cell.value = '背景概要'
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='left', vertical='center')
    row += 1

    ws.merge_cells(f'A{row}:B{row}')
    cell = ws[f'A{row}']
    cell.value = persona.get('background_summary', '')
    cell.font = content_font
    cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    ws.row_dimensions[row].height = 60
    row += 2

    # 核心议题
    ws.merge_cells(f'A{row}:B{row}')
    cell = ws[f'A{row}']
    cell.value = '核心议题'
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='left', vertical='center')
    row += 1

    for issue in persona.get('core_issues', []):
        ws.merge_cells(f'A{row}:B{row}')
        cell = ws[f'A{row}']
        cell.value = f'• {issue}'
        cell.font = content_font
        cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        row += 1

def _write_dialogue_sheet(ws, session_data):
    """填写Sheet2：咨询过程详细分析（10列）"""
    dialogue = session_data['dialogue']

    # 样式定义
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    content_font = Font(name='微软雅黑', size=9)
    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )

    # 设置列宽
    columns = [
        ('A', 6),   # 轮次
        ('B', 8),   # 角色
        ('C', 30),  # 对话内容
        ('D', 15),  # 技术使用
        ('E', 8),   # 评分
        ('F', 25),  # 评价
        ('G', 25),  # 更好的说法
        ('H', 20),  # 来访者情绪
        ('I', 20),  # 咨询师意图
        ('J', 30),  # 理论依据
    ]
    for col, width in columns:
        ws.column_dimensions[col].width = width

    # 表头
    headers = ['轮次', '角色', '对话内容', '技术使用', '评分', '评价', '更好的说法',
               '来访者情绪', '咨询师意图', '理论依据']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[1].height = 30

    # 数据行
    row = 2
    for turn in dialogue:
        role = '来访者' if turn['role'] == 'client' else '咨询师'
        content = turn['content']
        feedback = turn.get('supervision_feedback')

        # 写入基本信息
        ws.cell(row=row, column=1).value = turn['seq']
        ws.cell(row=row, column=2).value = role
        ws.cell(row=row, column=3).value = content

        # 如果有督导反馈，写入督导信息
        if feedback and isinstance(feedback, dict):
            ws.cell(row=row, column=4).value = feedback.get('technique_used', '')
            ws.cell(row=row, column=5).value = feedback.get('quality_score', '')
            ws.cell(row=row, column=6).value = feedback.get('comment', '')
            ws.cell(row=row, column=7).value = feedback.get('better_example', '')
            ws.cell(row=row, column=8).value = feedback.get('client_emotion', '')
            ws.cell(row=row, column=9).value = feedback.get('counselor_motivation', '')
            ws.cell(row=row, column=10).value = feedback.get('theoretical_basis', '')

        # 应用样式
        for col_idx in range(1, 11):
            cell = ws.cell(row=row, column=col_idx)
            cell.font = content_font
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            cell.border = thin_border

        ws.row_dimensions[row].height = max(40, len(content) // 30 * 15)
        row += 1

def _add_row(ws, row, label, value, header_font, header_fill, content_font):
    """添加一行数据（标签+值）"""
    cell_a = ws[f'A{row}']
    cell_a.value = label
    cell_a.font = header_font
    cell_a.fill = header_fill
    cell_a.alignment = Alignment(horizontal='left', vertical='center')

    cell_b = ws[f'B{row}']
    cell_b.value = value
    cell_b.font = content_font
    cell_b.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
