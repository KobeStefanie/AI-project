#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
完整读取4-8月所有数据
"""
import openpyxl

excel_file = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'
wb = openpyxl.load_workbook(excel_file, data_only=True)

print("=" * 80)
print("Excel中的所有Sheet：")
print("=" * 80)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"{sheet_name}: {ws.max_row}行 x {ws.max_column}列")

print("\n" + "=" * 80)
print("读取4-8月完整数据")
print("=" * 80)

months_data = {}

for month in ['4月', '5月', '6月', '7月', '8月']:
    if month not in wb.sheetnames:
        print(f"警告：{month}不存在")
        continue

    ws = wb[month]
    print(f"\n{month}表：")
    print(f"  总行数: {ws.max_row}")
    print(f"  总列数: {ws.max_column}")

    # 读取所有品种
    varieties = []
    for row_idx in range(3, ws.max_row + 1):
        name = ws.cell(row_idx, 3).value  # Col3: 物料名称
        if not name or name == '合计':
            continue

        # 读取关键字段
        item = {
            '序号': ws.cell(row_idx, 1).value,
            '车间': ws.cell(row_idx, 2).value,
            '物料名称': name,
            '类型': ws.cell(row_idx, 4).value,
            '实收数量': ws.cell(row_idx, 9).value,
            '原材料金额': ws.cell(row_idx, 47).value,
            '人工小计': ws.cell(row_idx, 54).value,
            '制造费用小计': ws.cell(row_idx, 65).value,
            '实际成本合计': ws.cell(row_idx, 67).value,
            '实际单位成本': ws.cell(row_idx, 69).value,
            '变动率': ws.cell(row_idx, 70).value,
        }

        varieties.append(item)

    months_data[month] = varieties
    print(f"  品种数: {len(varieties)}")

    # 计算月度汇总
    total_material = sum(float(v['原材料金额'] or 0) for v in varieties)
    total_labor = sum(float(v['人工小计'] or 0) for v in varieties)
    total_mfg = sum(float(v['制造费用小计'] or 0) for v in varieties)
    total_cost = sum(float(v['实际成本合计'] or 0) for v in varieties)

    print(f"  总成本: {total_cost/10000:.2f}万元")
    print(f"    原材料: {total_material/10000:.2f}万元 ({total_material/total_cost*100 if total_cost > 0 else 0:.1f}%)")
    print(f"    人工: {total_labor/10000:.2f}万元 ({total_labor/total_cost*100 if total_cost > 0 else 0:.1f}%)")
    print(f"    制造费用: {total_mfg/10000:.2f}万元 ({total_mfg/total_cost*100 if total_cost > 0 else 0:.1f}%)")

# 输出到文件
output_file = r'd:\AI-项目\12-成本控制\months_4_8_data.txt'
with open(output_file, 'w', encoding='utf-8') as f:
    for month, varieties in months_data.items():
        f.write(f"\n{'='*80}\n")
        f.write(f"{month}数据（{len(varieties)}个品种）\n")
        f.write('='*80 + '\n\n')

        for v in varieties[:5]:  # 只写前5个品种
            f.write(f"品种: {v['物料名称']}\n")
            f.write(f"  类型: {v['类型']}\n")
            f.write(f"  实收数量: {v['实收数量']}\n")
            f.write(f"  原材料: {v['原材料金额']}\n")
            f.write(f"  人工: {v['人工小计']}\n")
            f.write(f"  制造费用: {v['制造费用小计']}\n")
            f.write(f"  实际成本: {v['实际成本合计']}\n")
            f.write(f"  变动率: {v['变动率']}\n\n")

print(f"\n详细数据已写入: {output_file}")

wb.close()
