# -*- coding: utf-8 -*-
import os,json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
ROOT=r'D:\AI-项目\12-成本控制'; OUT=os.path.join(ROOT,'01-交付物归档','2026-09','织造车间成本体系建设_最新数据分析版_202609.pptx')
J=json.load(open(os.path.join(ROOT,'latest_cost_analysis.json'),encoding='utf8')); monthly=J['monthly']; products=J['products']; notes=J['notes']; ab=J['abnormal_money']; useab=J['use_abnormal']; months=J['months'];
NAVY=RGBColor(20,38,63); INK=RGBColor(35,45,56); MUT=RGBColor(103,116,128); TEAL=RGBColor(19,139,143); CORAL=RGBColor(222,91,75); GOLD=RGBColor(224,164,58); CREAM=RGBColor(247,245,239); WHITE=RGBColor(255,255,255); LINE=RGBColor(220,225,228); GREEN=RGBColor(69,145,95)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
def bg(s,c=CREAM): s.background.fill.solid(); s.background.fill.fore_color.rgb=c
def box(s,x,y,w,h,c,l=None,r=False): sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.color.rgb=l or c; return sh
def tx(s,z,x,y,w,h,sz=14,c=INK,b=False,a=PP_ALIGN.LEFT): q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); q.text_frame.word_wrap=True; p=q.text_frame.paragraphs[0]; p.text=z; p.alignment=a; p.font.name='Microsoft YaHei'; p.font.size=Pt(sz); p.font.bold=b; p.font.color.rgb=c; return q
def title(s,k,h,n=''): tx(s,k,.62,.35,3,.25,10,TEAL,True); tx(s,h,.62,.67,12,.55,27,NAVY,True); tx(s,n,.64,1.2,12,.3,11,MUT); box(s,.62,1.55,12.08,.02,LINE)
def foot(s,z='数据来源：最新成本核算汇总表（4-7月）；金额单位：元'): tx(s,z,.62,7.18,12,.18,8.5,MUT)
def kpi(s,x,y,w,l,v,n,c): box(s,x,y,w,1.1,WHITE,LINE,True); box(s,x,y,.08,1.1,c,c); tx(s,l,x+.2,y+.15,w-.3,.22,10,MUT,True); tx(s,v,x+.2,y+.42,w-.3,.38,24,NAVY,True); tx(s,n,x+.2,y+.84,w-.3,.18,9.5,MUT)
def bullet(s,z,x,y,w,c): box(s,x,y+.1,.1,.1,c,c,True); tx(s,z,x+.22,y,w-.22,.32,14,INK)
def table(s,x,y,headers,rows,widths,rowh=.4,headc=NAVY):
 xx=x
 for h,w in zip(headers,widths): box(s,xx,y,w,rowh,headc,headc); tx(s,h,xx+.04,y+.09,w-.08,.2,10,WHITE,True,PP_ALIGN.CENTER); xx+=w
 for i,row in enumerate(rows):
  xx=x; yy=y+rowh*(i+1); fill=WHITE if i%2==0 else RGBColor(242,245,246)
  for val,w in zip(row,widths): box(s,xx,yy,w,rowh,fill,LINE); tx(s,str(val),xx+.04,yy+.09,w-.08,.2,10,INK,False,PP_ALIGN.CENTER); xx+=w
# derived
actual_total=sum(monthly[m]['actual_cost'] for m in months); std_total=sum(monthly[m]['std_cost'] for m in months); var_total=actual_total-std_total
mat_total=sum(monthly[m]['actual_mat'] for m in months); qty_total=sum(monthly[m]['qty'] for m in months)
june=monthly['6月']; july=monthly['7月'];
# 1 cover
s=prs.slides.add_slide(blank); bg(s,NAVY); box(s,0,0,13.33,7.5,NAVY); box(s,0,0,.24,7.5,CORAL); tx(s,'织造车间',.85,1.15,3,.3,15,RGBColor(152,211,209),True); tx(s,'成本管理体系建设',.82,1.62,9,.7,40,WHITE,True); tx(s,'4-7月最新数据差异分析',.86,2.55,8,.4,24,RGBColor(222,229,234)); box(s,.86,3.42,5.2,.04,TEAL); tx(s,'从“总成本变化”追到“超支发生在哪里”',.86,3.75,8,.4,18,WHITE,True); tx(s,'财务部｜2026年9月',.86,6.78,3,.2,10,RGBColor(177,193,202))
# 2 conclusion
s=prs.slides.add_slide(blank); bg(s); title(s,'01｜管理层结论','总体低于标准，不代表没有问题：6月出现集中性材料超耗','新版数据已具备标准成本、实际用量和异常备注，可以把问题落到具体品种。'); kpi(s,.62,1.88,2.8,'4-7月实际成本',f'{actual_total/10000:.1f} 万元','累计生产量 116.6 吨',CORAL); kpi(s,3.62,1.88,2.8,'相对标准成本',f'{var_total/10000:+.1f} 万元',f'{var_total/std_total*100:+.1f}%，4-5月节约抵消6-7月超支',TEAL); kpi(s,6.62,1.88,2.8,'6月净超支',f'{june["variance"]/10000:.1f} 万元',f'单位成本 {june["actual_unit"]:.2f} 元/KG',CORAL); kpi(s,9.62,1.88,2.8,'备注异常',f'{J["counts"]["notes"]} 条','另有12条用量偏差>10%',GOLD); box(s,.62,3.35,12.08,2.65,WHITE,LINE,True); tx(s,'最重要的三个判断',.92,3.68,2.2,.25,12,TEAL,True); bullet(s,'6月不是普遍性超支，主要是 26S 棉双纱、26S 棉平纹、24S 小圆筒三类异常。',.94,4.1,11,CORAL); bullet(s,'三类异常在6月合计造成约131.6万元标准成本超支，超过当月净超支112.0万元。',.94,4.62,11,CORAL); bullet(s,'下一步应先做在产品盘点和领料追溯，再讨论采购价格或工艺优化。',.94,5.14,11,TEAL); foot(s)
# 3 monthly bridge
s=prs.slides.add_slide(blank); bg(s); title(s,'02｜月度标准成本差异','4-5月节约，6月转为超支，7月接近标准但异常仍未消失','把“净差异”与“异常记录”同时看，避免被平均数掩盖。'); rows=[]
for m in months:
 d=monthly[m]; rows.append([m,d['records'],f'{d["qty"]/1000:.1f}',f'{d["std_cost"]/10000:.1f}',f'{d["actual_cost"]/10000:.1f}',f'{d["variance"]/10000:+.1f}',f'{d["variance_pct"]:+.1f}%',f'{d["use_diff_pct"]:+.1f}%'])
table(s,.55,2.0,['月份','记录','产量(吨)','标准成本','实际成本','差异(万)','差异率','用纱差异'],rows,[.7,.65,1.0,1.2,1.2,1.0,1.0,1.1],.45)
box(s,.55,5.05,12.2,1.45,RGBColor(240,247,246),RGBColor(197,225,222),True); tx(s,'差异解释',.85,5.3,1.4,.25,13,TEAL,True); tx(s,f'4-5月累计节约 {(monthly["4月"]["variance"]+monthly["5月"]["variance"])/10000:.1f} 万元；6-7月累计超支 {(monthly["6月"]["variance"]+monthly["7月"]["variance"])/10000:.1f} 万元。6月实际用纱较标准高 {june["use_diff_pct"]:.1f}%，7月回落至 {july["use_diff_pct"]:.1f}%，但异常品种仍需单独追踪。',.85,5.65,11.5,.55,13,INK); foot(s)
# 4 abnormal root causes
s=prs.slides.add_slide(blank); bg(s); title(s,'03｜6月超支的根因线索','异常集中在“实际出库远超标准用纱”，备注指向在产品盘点或损耗问题','这里先呈现证据和管理假设，责任归因需要现场复核。'); ab6=[x for x in ab if x['month']=='6月'][:5]; rows=[]
for x in ab6: rows.append([x['name'],f'{x["qty"]:.0f}',f'{x["variance_amt"]/10000:.1f}',f'{x["variance_rate"]*100:.1f}%',f'{x["use_diff_pct"]:.1f}%',('在产品盘点/损耗待查' if x['note'] else '需补充原因')])
table(s,.65,1.95,['品种','产量KG','超支(万)','成本偏差','用纱偏差','原因线索'],rows,[2.5,1.0,1.1,1.1,1.1,3.4],.52)
box(s,.65,5.25,12.0,1.2,RGBColor(255,246,229),RGBColor(242,218,169),True); tx(s,'控制判断',.95,5.48,1.4,.25,13,CORAL,True); tx(s,'26S 棉双纱与26S 棉平纹的实际用纱均约为标准的2.06倍；优先复核：期初在产品数量、当月领料单、剩余在产品、次品和织损记录。',.95,5.84,11.3,.4,13,INK); foot(s)
# 5 focus portfolio
s=prs.slides.add_slide(blank); bg(s); title(s,'04｜重点品种排序','按“超支金额 × 可控性”排序，先处理异常金额最大的品种','单位成本偏差很高但金额很小的品种，先列入数据清理，不作为第一降本项目。'); rows=[]
for x in ab[:8]: rows.append([x['month'],x['name'],f'{x["variance_amt"]/10000:.1f}',f'{x["variance_rate"]*100:.1f}%',f'{x["use_diff_pct"]:.1f}%',('高' if x['variance_amt']>=300000 else '中')])
table(s,.7,2.0,['月份','异常品种','超支(万)','成本偏差','用纱偏差','优先级'],rows,[.8,2.7,1.1,1.1,1.1,1.0],.48)
box(s,8.45,2.0,4.2,3.85,WHITE,LINE,True); tx(s,'处理顺序',8.8,2.3,2,.3,15,NAVY,True); bullet(s,'P0：26S棉平纹（6月+53.9万）',8.8,3.0,3.2,CORAL); bullet(s,'P0：26S棉双纱（6月+53.4万，7月+41.6万）',8.8,3.62,3.2,CORAL); bullet(s,'P1：24S小圆筒（6月+24.3万，7月+12.0万）',8.8,4.24,3.2,GOLD); bullet(s,'P2：40D1*1罗纹，金额大但偏差约1%-4%，纳入常规监控。',8.8,4.86,3.2,TEAL); foot(s)
# 6 data quality
s=prs.slides.add_slide(blank); bg(s); title(s,'05｜数据质量与标准成本','数据已经能做差异分析，但标准用量和在产品盘点仍是关键控制点','异常备注数量下降不等于问题消失，必须让异常有闭环。'); kpi(s,.7,1.95,2.8,'实际用量偏差>10%',f'{J["counts"]["use_gt10"]} 条','需要复核领料和盘点',CORAL); kpi(s,3.7,1.95,2.8,'实际成本低于标准',f'{J["counts"]["negative_variance_records"]} 条','节约数据也要确认是否漏计',TEAL); kpi(s,6.7,1.95,2.8,'实际成本高于标准',f'{J["counts"]["positive_variance_records"]} 条','重点看金额而非只看比例',GOLD); kpi(s,9.7,1.95,2.8,'缺少实际成本',f'{J["counts"]["actual_zero"]} 条','不纳入经营结论',MUT); box(s,.7,3.55,12,2.75,WHITE,LINE,True); tx(s,'建议建立异常闭环字段',1.0,3.87,3,.3,15,NAVY,True); bullet(s,'异常类型：盘点差异 / 领料漏记 / 工艺标准错误 / 真实织损 / 产品切换余料。',1.0,4.45,10.8,TEAL); bullet(s,'责任人：仓库、车间、技术、财务分别确认，不用“数据异常”笼统结案。',1.0,4.92,10.8,CORAL); bullet(s,'截止时间：月结后T+2完成复核，T+5完成成本报告。',1.0,5.39,10.8,GOLD); foot(s)
# 7 action
s=prs.slides.add_slide(blank); bg(s); title(s,'06｜9月行动方案','把异常分析变成现场动作：先盘点、再追溯、后更新标准','每个动作都对应最新 Excel 中已有的字段。'); rows=[['1','复核26S/24S','核对期初在产品、领料单、剩余在产品、次品与织损','6月异常金额全部有结论'],['2','固定领料链路','订单→领料→胚布入库→在产品盘点→月结','实际用量偏差>10%必须解释'],['3','维护标准成本','工艺单标准用纱、织损率、纱线单价按变更留痕','标准成本月度可复核'],['4','形成差异报告','按金额排序，区分价格/用量/盘点/工艺四类','T+5提交管理层']]; table(s,.7,1.95,['步骤','重点','动作','验收'],rows,[.6,1.7,6.7,2.3],.62); box(s,.7,5.8,12,.55,RGBColor(240,247,246),RGBColor(197,225,222),True); tx(s,'降本测算暂不承诺年度金额：先把6月异常的131.6万元拆清楚，再计算可兑现空间。',.95,5.98,11.4,.2,12,INK,True); foot(s)
# 8 asks
s=prs.slides.add_slide(blank); bg(s,NAVY); tx(s,'07｜需要管理层确认',.78,.62,4,.3,12,RGBColor(150,219,216),True); tx(s,'先把材料差异查清楚，再把降本目标定下来',.78,1.12,11,.6,29,WHITE,True); asks=[('01','批准专项盘点','优先复核26S棉双纱、26S棉平纹、24S小圆筒。',CORAL),('02','明确跨部门责任','仓库、车间、技术、财务共同关闭异常。',GOLD),('03','固定月结节奏','T+2复核，T+5提交成本差异报告。',TEAL)]
for i,(n,h,b,c) in enumerate(asks): y=2.4+i*1.15; box(s,.82,y,.55,.55,c,c,True); tx(s,n,.82,y+.14,.55,.2,13,WHITE,True,PP_ALIGN.CENTER); tx(s,h,1.65,y+.02,3,.28,18,WHITE,True); tx(s,b,1.65,y+.38,9.9,.3,13,RGBColor(218,229,235))
box(s,.82,6.08,11.65,.62,RGBColor(30,57,83),RGBColor(74,112,137),True); tx(s,'本版结论：4-5月整体节约，6月异常超支集中在用纱和盘点，7月仍需持续追踪。',1.08,6.28,11.1,.22,16,RGBColor(237,243,245),True,PP_ALIGN.CENTER); tx(s,'财务部｜2026年9月',.82,7.12,3,.18,9,RGBColor(153,176,188))
os.makedirs(os.path.dirname(OUT),exist_ok=True); prs.save(OUT); print(OUT)
