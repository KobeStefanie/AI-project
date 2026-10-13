# -*- coding: utf-8 -*-
import os,json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
ROOT=r'D:\AI-项目\12-成本控制'; OUT=os.path.join(ROOT,'01-交付物归档','2026-09','织造车间成本体系建设_成本分析版_202609.pptx')
J=json.load(open(os.path.join(ROOT,'deep_cost_analysis.json'),encoding='utf8')); monthly=J['monthly']; valid=J['valid_7_8']; prods=J['products']; conc=J['concentration']; miss=J['missing_material']; small=J['small_qty']; rec=J['record_count']
NAVY=RGBColor(20,38,63); INK=RGBColor(35,45,56); MUT=RGBColor(103,116,128); TEAL=RGBColor(19,139,143); CORAL=RGBColor(222,91,75); GOLD=RGBColor(224,164,58); CREAM=RGBColor(247,245,239); WHITE=RGBColor(255,255,255); LINE=RGBColor(220,225,228); GREEN=RGBColor(69,145,95)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
def bg(s,c=CREAM): s.background.fill.solid(); s.background.fill.fore_color.rgb=c
def box(s,x,y,w,h,c,l=None,r=False): sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.color.rgb=l or c; return sh
def tx(s,z,x,y,w,h,sz=14,c=INK,b=False,a=PP_ALIGN.LEFT): q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); q.text_frame.word_wrap=True; p=q.text_frame.paragraphs[0]; p.text=z; p.alignment=a; p.font.name='Microsoft YaHei'; p.font.size=Pt(sz); p.font.bold=b; p.font.color.rgb=c; return q
def title(s,k,h,n=''): tx(s,k,.62,.35,3,.25,10,TEAL,True); tx(s,h,.62,.67,12,.55,27,NAVY,True); tx(s,n,.64,1.2,12,.3,11,MUT); box(s,.62,1.55,12.08,.02,LINE)
def foot(s,z='数据来源：4-8月成本核算汇总表；金额单位：元'): tx(s,z,.62,7.18,12,.18,8.5,MUT)
def kpi(s,x,y,w,l,v,n,c): box(s,x,y,w,1.1,WHITE,LINE,True); box(s,x,y,.08,1.1,c,c); tx(s,l,x+.2,y+.15,w-.3,.22,10,MUT,True); tx(s,v,x+.2,y+.42,w-.3,.38,24,NAVY,True); tx(s,n,x+.2,y+.84,w-.3,.18,9.5,MUT)
def bullet(s,z,x,y,w,c): box(s,x,y+.1,.1,.1,c,c,True); tx(s,z,x+.22,y,w-.22,.32,14,INK)
def table(s,x,y,headers,rows,widths,rowh=.38,headc=NAVY):
 xx=x
 for h,w in zip(headers,widths): box(s,xx,y,w,rowh,headc,headc); tx(s,h,xx+.05,y+.08,w-.1,.2,10,WHITE,True,PP_ALIGN.CENTER); xx+=w
 for i,row in enumerate(rows):
  xx=x; yy=y+rowh*(i+1); fill=WHITE if i%2==0 else RGBColor(242,245,246)
  for val,w in zip(row,widths): box(s,xx,yy,w,rowh,fill,LINE); tx(s,str(val),xx+.05,yy+.08,w-.1,.2,10,INK,False,PP_ALIGN.CENTER); xx+=w
# 1
s=prs.slides.add_slide(blank); bg(s,NAVY); box(s,0,0,13.33,7.5,NAVY); box(s,0,0,.24,7.5,CORAL); tx(s,'织造车间',.85,1.15,3,.3,15,RGBColor(152,211,209),True); tx(s,'成本管理体系建设',.82,1.62,9,.7,40,WHITE,True); tx(s,'4-8月成本数据深度分析版',.86,2.55,8,.4,24,RGBColor(222,229,234)); box(s,.86,3.42,5.2,.04,TEAL); tx(s,'从“成本很高”走到“知道先管哪一项”',.86,3.75,8,.4,18,WHITE,True); tx(s,'财务部｜2026年9月',.86,6.78,3,.2,10,RGBColor(177,193,202))
# 2
s=prs.slides.add_slide(blank); bg(s); title(s,'01｜管理层结论','成本问题不是平均分布，而是集中在少数主力品种和原材料驱动上','本页只保留可直接支持管理动作的结论。'); kpi(s,.62,1.88,2.8,'累计实际成本','7,135 万元','4-8月全部记录',CORAL); kpi(s,3.62,1.88,2.8,'前2品种成本占比','56.2%','华棉118D + 40D1*1罗纹',TEAL); kpi(s,6.62,1.88,2.8,'原材料占比','89%-94%','每月实际成本构成',GOLD); kpi(s,9.62,1.88,2.8,'6月单位成本峰值','590.42 元/KG','较5月上升28.0%',GREEN); box(s,.62,3.35,12.08,2.65,WHITE,LINE,True); tx(s,'三条经营判断',.92,3.68,2,.25,12,TEAL,True); bullet(s,'第一抓手：先盯原材料单价、用量和织损，人工与制造费用不是当前主要矛盾。',.94,4.1,11,CORAL); bullet(s,'第一批对象：华棉118D和40D1*1罗纹合计占成本56.2%，优先建立两类品种的材料用量基准。',.94,4.62,11,TEAL); bullet(s,'第一项补数：现有数据不能区分材料价格变化和材料用量变化，必须补充纱线单价、标准用量和领料量。',.94,5.14,11,GOLD); foot(s)
# 3 monthly
s=prs.slides.add_slide(blank); bg(s); title(s,'02｜月度成本变化','6月成本峰值由材料单位成本跳升驱动，8月下降也主要来自材料单位成本回落','总成本变化同时受产量影响，因此用单位成本和成本/公斤拆解。'); rows=[]
for m in ['4月','5月','6月','7月','8月']:
 d=monthly[m]; rows.append([m,d['n'],f'{d["qty"]/1000:.1f}',f'{d["cost"]/10000:.1f}',f'{d["unit"]:.2f}',f'{d["material_per_kg"]:.2f}',f'{d["labor_per_kg"]:.2f}',f'{d["mfg_per_kg"]:.2f}'])
table(s,.7,2.0,['月份','品种','产量(吨)','成本(万元)','单位成本','材料/公斤','人工/公斤','制造/公斤'],rows,[.8,.75,1.0,1.15,1.15,1.25,1.15,1.35],.42)
box(s,.7,5.0,12.0,1.55,RGBColor(240,247,246),RGBColor(197,225,222),True); tx(s,'变化拆解',1.0,5.25,1.5,.25,13,TEAL,True); tx(s,'5→6月：单位成本 +28.0%，材料/公斤 +34.4%，是主要驱动；6→7月：材料/公斤 -10.9%，但产量上升，单位成本仍高；7→8月：单位成本 -33.0%，材料/公斤 -35.8%，需结合采购单价和产品结构验证原因。',1.0,5.62,11.3,.65,13,INK); foot(s)
# 4 concentration
s=prs.slides.add_slide(blank); bg(s); title(s,'03｜成本集中度','成本高度集中：抓住前5个品种，就抓住约75%的成本池','这决定了改善顺序应先做重点品种，而不是平均分配精力。'); rows=[]; cum=0
for i,p in enumerate(prods[:10],1): cum+=p['cost_share']; rows.append([i,p['name'],f'{p["cost"]/10000:.1f}',f'{p["cost_share"]:.1f}%',f'{cum:.1f}%',f'{p["unit"]:.2f}'])
table(s,.8,1.95,['排名','品种','成本(万元)','占比','累计','单位成本'],rows,[.6,3.1,1.4,1.0,1.0,1.3],.42)
box(s,8.5,2.0,4.0,3.95,WHITE,LINE,True); tx(s,'管理含义',8.85,2.32,2,.3,15,NAVY,True); tx(s,'56.2%',8.85,2.85,2.4,.55,34,CORAL,True); tx(s,'前2个品种成本占比',8.88,3.45,2.8,.25,11,MUT); bullet(s,'华棉118D：成本2,007万元，单位成本420.53元/KG。',8.85,4.0,3.2,CORAL); bullet(s,'40D1*1罗纹：成本2,004万元，单位成本642.40元/KG。',8.85,4.63,3.2,TEAL); bullet(s,'先建立两品种的材料价格、标准用量、实际领料三张对照表。',8.85,5.26,3.2,GOLD); foot(s)
# 5 products
s=prs.slides.add_slide(blank); bg(s); title(s,'04｜重点品种诊断','高成本品种不等于高损耗品种：先区分“结构贵”与“管理异常”','当前数据支持优先级判断，暂不支持直接判定责任归因。'); rows=[]
for p in prods[:8]: rows.append([p['name'],f'{p["qty"]/1000:.1f}',f'{p["unit"]:.2f}',f'{p["material_pct"]:.1f}%',f'{p["material_per_kg"]:.2f}',f'{p["cost"]/10000:.1f}'])
table(s,.7,1.95,['品种','产量(吨)','单位成本','材料占比','材料/公斤','成本(万元)'],rows,[2.7,1.0,1.2,1.2,1.35,1.2],.48)
box(s,.7,6.0,12.0,.62,RGBColor(255,246,229),RGBColor(242,218,169),True); tx(s,'判断边界：26S棉平纹单位成本1,189.10元/KG，但累计成本仅131万元；不能仅凭单位成本把它列为第一降本对象。',.95,6.19,11.5,.2,11,INK,True); foot(s)
# 6 data vs causality
s=prs.slides.add_slide(blank); bg(s); title(s,'05｜数据能回答什么','现在能确认“哪里异常”，还不能确认“异常由谁造成”','把可证实事实和待验证假设分开，避免把数据缺口误报成经营结论。'); box(s,.75,1.95,5.8,4.7,RGBColor(240,247,246),RGBColor(197,225,222),True); tx(s,'已证实',1.1,2.28,2,.3,16,TEAL,True); bullet(s,'6月材料/公斤552.01元，为5个月最高。',1.1,3.0,4.8,TEAL); bullet(s,'8月材料/公斤315.87元，较7月下降35.8%。',1.1,3.65,4.8,TEAL); bullet(s,'前2品种占成本56.2%，改善资源应集中。',1.1,4.3,4.8,TEAL); bullet(s,'69条原材料金额为0，成本可追溯性不足。',1.1,4.95,4.8,TEAL); box(s,6.8,1.95,5.8,4.7,RGBColor(255,246,229),RGBColor(242,218,169),True); tx(s,'待验证',7.15,2.28,2,.3,16,CORAL,True); bullet(s,'材料/公斤变化来自采购价格还是实际用量？',7.15,3.0,4.8,CORAL); bullet(s,'8月成本下降是否来自产品结构变化？',7.15,3.65,4.8,CORAL); bullet(s,'高单位成本是否由小批量分摊造成？',7.15,4.3,4.8,CORAL); bullet(s,'实际领料是否超过工艺单标准用量？',7.15,4.95,4.8,CORAL); foot(s)
# 7 action quantified
s=prs.slides.add_slide(blank); bg(s); title(s,'06｜改善动作与量化目标','把分析变成可追踪的成本控制闭环','每项动作都要有数据字段、责任人和验收口径。'); rows=[['9月','补齐材料链路','主力品种绑定订单；记录纱线单价/领料量/标准用量','完整率≥90%'],['10月','做材料差异桥','价格差异、用量差异、织损差异分开','每月可解释'],['11月','做重点品种预警','前5品种单位材料成本环比>10%自动复核','异常T+2反馈']]; table(s,.75,2.0,['时间','动作','落地内容','验收'],rows,[1.0,1.7,6.2,2.2],.65); box(s,.75,5.0,12,1.5,WHITE,LINE,True); tx(s,'降本测算原则',1.05,5.28,2,.25,14,NAVY,True); tx(s,'先不承诺“年降本金额”。等材料价格、标准用量和实际领料连续运行1个月后，再用前5品种的材料差异计算可兑现的降本空间。',1.05,5.7,11.2,.45,14,INK); foot(s)
# 8 asks
s=prs.slides.add_slide(blank); bg(s,NAVY); tx(s,'07｜需要管理层确认',.78,.62,4,.3,12,RGBColor(150,219,216),True); tx(s,'下一步不是再做一张表，而是让数据在现场产生',.78,1.12,11,.6,29,WHITE,True); asks=[('01','按订单领料','生产部配合，先覆盖华棉118D和40D1*1罗纹。',CORAL),('02','材料差异三分法','财务、采购、生产共同确认价格/用量/织损字段。',GOLD),('03','月度成本会','固定复盘前5品种，异常项必须有责任人与动作。',TEAL)]
for i,(n,h,b,c) in enumerate(asks): y=2.4+i*1.15; box(s,.82,y,.55,.55,c,c,True); tx(s,n,.82,y+.14,.55,.2,13,WHITE,True,PP_ALIGN.CENTER); tx(s,h,1.65,y+.02,3,.28,18,WHITE,True); tx(s,b,1.65,y+.38,9.9,.3,13,RGBColor(218,229,235))
box(s,.82,6.08,11.65,.62,RGBColor(30,57,83),RGBColor(74,112,137),True); tx(s,'本版结论：先抓前2个品种，先拆材料差异，先把领料数据做实。',1.08,6.28,11.1,.22,16,RGBColor(237,243,245),True,PP_ALIGN.CENTER); tx(s,'财务部｜2026年9月',.82,7.12,3,.18,9,RGBColor(153,176,188))
os.makedirs(os.path.dirname(OUT),exist_ok=True); prs.save(OUT); print(OUT)
