#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
检查Excel文件的完整结构和数据量
"""
import openpyxl

excel_path = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'

# 打开Excel文件（只读数据，不读公式）
wb = openpyxl.load_workbook(excel_path, data_only=True)

output = []
output.append("=" * 80)
output.append("Excel文件完整结构分析")
output.append("=" * 80)
output.append("")

# 1. 列出所有sheet及其大小
output.append(f"总共有 {len(wb.sheetnames)} 个sheet：")
output.append("")

total_cells = 0
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    cells = ws.max_row * ws.max_column
    total_cells += cells
    output.append(f"  Sheet: {sheet_name}")
    output.append(f"    行数: {ws.max_row}")
    output.append(f"    列数: {ws.max_column}")
    output.append(f"    单元格数: {cells:,}")
    output.append("")

output.append("=" * 80)
output.append(f"总单元格数: {total_cells:,}")
output.append("")

# 2. 检查8月sheet的数据密度（有多少单元格有值）
ws_aug = wb['8月']
non_empty = 0
for row in ws_aug.iter_rows():
    for cell in row:
        if cell.value is not None:
            non_empty += 1

output.append(f"8月sheet数据密度: {non_empty:,} / {ws_aug.max_row * ws_aug.max_column:,} ({non_empty/(ws_aug.max_row * ws_aug.max_column)*100:.1f}%)")
output.append("")

# 3. 看看8月sheet有没有多张表格（通过标题行识别）
output.append("=" * 80)
output.append("8月sheet可能的表格结构（检查前10行）：")
output.append("")

for row_idx in range(1, min(11, ws_aug.max_row + 1)):
    row_data = []
    for col_idx in range(1, min(20, ws_aug.max_column + 1)):
        val = ws_aug.cell(row_idx, col_idx).value
        if val:
            row_data.append(f"[C{col_idx}]{val}")

    if row_data:
        output.append(f"Row{row_idx}: {' | '.join(row_data[:5])}")

output.append("")
output.append("=" * 80)
output.append("结论：数据量评估")
output.append("")
output.append(f"如果只读8月sheet的有效数据（{non_empty:,}个单元格），")
output.append(f"预计加载到内存约 {non_empty * 50 / 1024 / 1024:.1f} MB")
output.append("（假设每个单元格平均50字节）")
output.append("")
output.append("建议：")
output.append("  - 如果数据 < 10MB：可以全部加载")
output.append("  - 如果数据 > 10MB：建议按需读取指定列")

# 写入文件
output_file = 'd:/AI-项目/12-成本控制/excel_full_structure.txt'
with open(output_file, 'w', encoding='utf-8') as f:
    f.write('\n'.join(output))

print(f"Done: {output_file}")
