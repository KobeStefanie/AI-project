#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
读取带公式的成本核算表，理解数据逻辑
"""
import openpyxl

excel_path = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算表-7月-2版.xlsx'

# 读取带公式的Excel
wb = openpyxl.load_workbook(excel_path, data_only=False)

output = []
output.append("=" * 80)
output.append("成本核算表-7月-2版 结构分析")
output.append("=" * 80)
output.append("")

# 1. 列出所有sheet
output.append(f"Sheet列表（共{len(wb.sheetnames)}个）：")
for i, sheet_name in enumerate(wb.sheetnames, 1):
    ws = wb[sheet_name]
    output.append(f"  {i}. {sheet_name} ({ws.max_row}行 × {ws.max_column}列)")
output.append("")

# 2. 如果有成本核算相关的sheet，读取其结构
for sheet_name in wb.sheetnames[:3]:  # 只看前3个sheet
    output.append("=" * 80)
    output.append(f"Sheet: {sheet_name}")
    output.append("=" * 80)

    ws = wb[sheet_name]

    # 读取前5行标题
    output.append("\n前5行标题：")
    for row_idx in range(1, min(6, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, min(15, ws.max_column + 1)):
            val = ws.cell(row_idx, col_idx).value
            if val:
                row_data.append(f"[C{col_idx}]{val}")
        if row_data:
            output.append(f"  Row{row_idx}: {' | '.join(row_data[:8])}")

    output.append("")

    # 检查是否有公式（抽查几个可能有公式的单元格）
    output.append("公式示例（Row3）：")
    for col_idx in [10, 15, 20, 25, 30]:
        if col_idx <= ws.max_column:
            cell = ws.cell(3, col_idx)
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                col_header = ws.cell(1, col_idx).value or ws.cell(2, col_idx).value or f"Col{col_idx}"
                output.append(f"  Col{col_idx} ({col_header}): {cell.value[:80]}")

    output.append("")

# 写入文件
output_file = 'd:/AI-项目/12-成本控制/formula_excel_structure.txt'
with open(output_file, 'w', encoding='utf-8') as f:
    f.write('\n'.join(output))

print(f"Done: {output_file}")
