#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
完整分析成本核算Excel的数据结构和逻辑
"""
import openpyxl
import sys
import os

def main():
    # 读取带公式的Excel
    formula_file = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算表-7月-2版.xlsx'

    print("=" * 80)
    print("第一部分：读取31个sheet的完整结构")
    print("=" * 80)

    wb = openpyxl.load_workbook(formula_file, data_only=False)

    # 1. 列出所有sheet
    output_lines = []
    output_lines.append("\n所有Sheet列表（共{}个）：".format(len(wb.sheetnames)))
    for i, name in enumerate(wb.sheetnames, 1):
        ws = wb[name]
        output_lines.append(f"{i:2d}. {name:30s} - {ws.max_row:6d}行 x {ws.max_column:3d}列")

    # 2. 重点分析核心sheet
    core_sheets = ['字典', '成本核算', '工艺单', '原纱汇总_库', '胚布入库_ERP',
                   '挡车工_织造', '验布工_织造', '财务_制造']

    for sheet_name in core_sheets:
        if sheet_name not in wb.sheetnames:
            continue

        output_lines.append("\n" + "=" * 80)
        output_lines.append(f"Sheet: {sheet_name}")
        output_lines.append("=" * 80)

        ws = wb[sheet_name]

        # 读取前5行的完整内容
        output_lines.append("\n前5行完整内容：")
        for row_idx in range(1, min(6, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, min(21, ws.max_column + 1)):  # 只读前20列避免太长
                cell = ws.cell(row_idx, col_idx)
                val = cell.value
                if val is None:
                    val = ""
                elif isinstance(val, str) and len(val) > 20:
                    val = val[:20] + "..."
                row_data.append(f"Col{col_idx}:{val}")
            output_lines.append(f"  Row{row_idx}: " + " | ".join(row_data))

        # 如果是成本核算sheet，特别分析公式列
        if sheet_name == '成本核算':
            output_lines.append("\n关键列的公式示例（第3行）：")
            key_cols = [9, 47, 54, 65, 67, 69, 70]  # 实收数量、原材料、人工、制造费用、实际成本、单位成本、变动率
            for col_idx in key_cols:
                cell = ws.cell(3, col_idx)
                header1 = ws.cell(1, col_idx).value or ""
                header2 = ws.cell(2, col_idx).value or ""
                output_lines.append(f"  Col{col_idx} ({header1}/{header2})")
                output_lines.append(f"    公式: {cell.value if isinstance(cell.value, str) and cell.value.startswith('=') else '(无公式)'}")
                output_lines.append(f"    数据: {wb[sheet_name].cell(3, col_idx).value if openpyxl.load_workbook(formula_file, data_only=True)[sheet_name].cell(3, col_idx).value else '空'}")

    # 写入输出文件
    output_file = r'd:\AI-项目\12-成本控制\excel_full_analysis.txt'
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write('\n'.join(output_lines))

    print(f"分析完成，结果写入: {output_file}")

    # 3. 读取汇总表的完整数据
    print("\n" + "=" * 80)
    print("第二部分：读取4-8月汇总表的完整数据")
    print("=" * 80)

    summary_file = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'
    wb2 = openpyxl.load_workbook(summary_file, data_only=True)

    summary_output = []

    for month in ['4月', '5月', '6月', '7月', '8月']:
        if month not in wb2.sheetnames:
            continue

        ws = wb2[month]
        summary_output.append(f"\n{'='*80}")
        summary_output.append(f"{month}表 - 完整列结构（71列）")
        summary_output.append('='*80)

        # 读取完整的双层表头
        summary_output.append("\n双层表头：")
        for col_idx in range(1, min(72, ws.max_column + 1)):
            h1 = ws.cell(1, col_idx).value or ""
            h2 = ws.cell(2, col_idx).value or ""
            summary_output.append(f"Col{col_idx:2d}: [{h1}] / [{h2}]")

        # 读取前5行数据
        summary_output.append(f"\n前5行数据：")
        for row_idx in range(3, min(8, ws.max_row + 1)):
            name = ws.cell(row_idx, 3).value  # 物料名称
            if not name:
                continue
            summary_output.append(f"\nRow{row_idx}: {name}")
            # 只读关键列
            key_cols = {
                3: "物料名称", 4: "加工品种", 5: "自产品种",
                9: "实收数量", 47: "原材料金额", 54: "人工小计",
                65: "制造费用小计", 67: "实际成本合计", 69: "实际单位成本", 70: "变动率"
            }
            for col_idx, col_name in key_cols.items():
                val = ws.cell(row_idx, col_idx).value
                summary_output.append(f"  Col{col_idx}({col_name}): {val}")

    # 写入汇总表分析
    summary_file_output = r'd:\AI-项目\12-成本控制\summary_full_analysis.txt'
    with open(summary_file_output, 'w', encoding='utf-8') as f:
        f.write('\n'.join(summary_output))

    print(f"汇总表分析完成，结果写入: {summary_file_output}")

if __name__ == '__main__':
    main()
