# -*- coding: utf-8 -*-
import openpyxl, json, os, math, statistics
path=r'C:\Users\Administrator\Desktop\华盛\成本核算汇总-标准成本 - 加工与自产.xlsx'
w=openpyxl.load_workbook(path,data_only=True,read_only=True)
print('SHEETS',w.sheetnames)
for sn in ['单位成本对比表','成本明细','单本明细','损耗对比','成本对比 ']:
 ws=w[sn]; print('SHEET',sn,'rows',ws.max_row,'cols',ws.max_column)
 names=[]
 for i,row in enumerate(ws.iter_rows(min_row=3,values_only=True),3):
  if row[2] is not None and str(row[2]).strip() and str(row[2]) not in ('合计','汇总'):
   names.append(str(row[2]))
 print('name_count',len(names),'first',names[:10])
for sn in ['4月','5月','6月','7月']:
 ws=w[sn]; print('MONTH',sn,'rows',ws.max_row,'cols',ws.max_column)
 # headers row0 row1
 h0=[x for x in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
 h1=[x for x in next(ws.iter_rows(min_row=2,max_row=2,values_only=True))]
 for idx,(a,b) in enumerate(zip(h0,h1),1):
  if a is not None or b is not None: print(idx,repr(a),repr(b))
 print('--- sample names/comments ---')
 for i,row in enumerate(ws.iter_rows(min_row=3,max_row=min(ws.max_row,15),values_only=True),3): print(i,row[2],row[8],row[15],row[16],row[44],row[46],row[71],row[73],row[74],row[75],row[76],row[78])
