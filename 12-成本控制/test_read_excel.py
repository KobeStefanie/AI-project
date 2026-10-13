#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
测试：读取Excel的不同方式
"""
import openpyxl
import pandas as pd

excel_file = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'

print("=" * 80)
print("方法1：openpyxl读取（data_only=True，读计算值）")
print("=" * 80)

wb1 = openpyxl.load_workbook(excel_file, data_only=True)
ws1 = wb1['8月']

print(f"Sheet名称：8月")
print(f"行数：{ws1.max_row}, 列数：{ws1.max_column}")
print(f"\n第3行（第一个品种）前10列的值：")
for col in range(1, 11):
    val = ws1.cell(3, col).value
    print(f"  Col{col}: {val}")

wb1.close()

print("\n" + "=" * 80)
print("方法2：pandas读取（默认读计算值）")
print("=" * 80)

df = pd.read_excel(excel_file, sheet_name='8月', header=None)
print(f"DataFrame形状：{df.shape}")
print(f"\n前3行，前10列：")
print(df.iloc[:3, :10])

print("\n" + "=" * 80)
print("方法3：openpyxl读取公式（data_only=False）")
print("=" * 80)

wb2 = openpyxl.load_workbook(excel_file, data_only=False)
ws2 = wb2['8月']

print(f"\n第3行的Col47（原材料金额）：")
cell = ws2.cell(3, 47)
print(f"  公式：{cell.value if isinstance(cell.value, str) and '=' in str(cell.value) else '(不是公式)'}")
print(f"  计算值：需要data_only=True才能看到")

wb2.close()

print("\n" + "=" * 80)
print("你的问题是：")
print("=" * 80)
print("1. 我读取的数据不对？（数值错误）")
print("2. 我没读取公式？（需要看公式逻辑）")
print("3. 我没读取全部列？（只读了部分列）")
print("4. 我没理解数据关系？（数据之间的逻辑）")
print("\n请告诉我具体是哪个问题。")
