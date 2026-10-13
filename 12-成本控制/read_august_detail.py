#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
读取8月数据的完整逻辑
"""
import openpyxl

summary_file = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'
wb = openpyxl.load_workbook(summary_file, data_only=True)

ws = wb['8月']

print("=" * 100)
print("8月成本数据 - 完整明细（前10个品种）")
print("=" * 100)

output = []

for row_idx in range(3, min(13, ws.max_row + 1)):
    name = ws.cell(row_idx, 3).value  # Col3: 物料名称
    if not name or name == '合计':
        continue

    output.append(f"\n{'='*100}")
    output.append(f"Row{row_idx}: {name}")
    output.append('='*100)

    # 基本信息
    output.append("\n【基本信息】")
    output.append(f"  车间: {ws.cell(row_idx, 2).value}")
    output.append(f"  类型: {ws.cell(row_idx, 4).value}")  # 加工/自产
    output.append(f"  实收数量: {ws.cell(row_idx, 9).value} {ws.cell(row_idx, 8).value}")
    output.append(f"  重量: {ws.cell(row_idx, 10).value}")

    # 原材料成本（3种纱支）
    output.append("\n【原材料成本】")
    for yarn_idx in [1, 2, 3]:
        base_col = 11 + (yarn_idx - 1) * 11  # 纱支1从Col11开始，每个纱支占11列
        yarn_code = ws.cell(row_idx, base_col).value  # 纱支编号
        if yarn_code:
            yarn_name = ws.cell(row_idx, base_col + 1).value
            ratio = ws.cell(row_idx, base_col + 2).value
            std_qty = ws.cell(row_idx, base_col + 5).value  # 标准用纱
            act_qty = ws.cell(row_idx, base_col + 6).value  # 实际出库
            unit_price = ws.cell(row_idx, base_col + 8).value  # 出库单价
            act_amount = ws.cell(row_idx, base_col + 10).value  # 实际出库金额

            output.append(f"  纱支{yarn_idx}: {yarn_name} ({yarn_code})")
            output.append(f"    比例: {ratio}%, 标准用量: {std_qty}, 实际出库: {act_qty}")
            output.append(f"    单价: {unit_price}, 实际金额: {act_amount}")

    output.append(f"  >>> 原材料合计: {ws.cell(row_idx, 47).value}")  # Col47: 实际金额小计

    # 人工成本
    output.append("\n【人工成本】")
    output.append(f"  挡车工: 比例{ws.cell(row_idx, 48).value}%, 工资{ws.cell(row_idx, 49).value}")
    output.append(f"  验布工: 比例{ws.cell(row_idx, 50).value}%, 工资{ws.cell(row_idx, 51).value}")
    output.append(f"  其他人员: 比例{ws.cell(row_idx, 52).value}%, 工资{ws.cell(row_idx, 53).value}")
    output.append(f"  >>> 人工小计: {ws.cell(row_idx, 54).value}")

    # 制造费用
    output.append("\n【制造费用】")
    output.append(f"  分配比例: {ws.cell(row_idx, 55).value}")
    output.append(f"  折旧: {ws.cell(row_idx, 56).value}")
    output.append(f"  机物料: {ws.cell(row_idx, 57).value}")
    output.append(f"  水电费: {ws.cell(row_idx, 59).value}")
    output.append(f"  其他: {ws.cell(row_idx, 62).value}")
    output.append(f"  >>> 制造费用小计: {ws.cell(row_idx, 65).value}")

    # 成本汇总
    output.append("\n【成本汇总】")
    output.append(f"  标准成本合计: {ws.cell(row_idx, 66).value}")
    output.append(f"  实际成本合计: {ws.cell(row_idx, 67).value}")
    output.append(f"  标准单位成本: {ws.cell(row_idx, 68).value}")
    output.append(f"  实际单位成本: {ws.cell(row_idx, 69).value}")
    output.append(f"  变动率: {ws.cell(row_idx, 70).value}")

# 输出到文件
output_file = r'd:\AI-项目\12-成本控制\august_detail.txt'
with open(output_file, 'w', encoding='utf-8') as f:
    f.write('\n'.join(output))

print('\n'.join(output))
print(f"\n详细数据已写入: {output_file}")

wb.close()
