# -*- coding: utf-8 -*-
import os,pickle
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
ROOT=r'D:\AI-项目\12-成本控制'; OUT=os.path.join(ROOT,'01-交付物归档','2026-09','织造车间成本体系建设_优化版_202609.pptx')
with open(os.path.join(ROOT,'months_data_correct.pkl'),'rb') as f: M=pickle.load(f)
Q='实收数量'; C='实际成本合计'; MAT='原材料金额'; LAB='人工小计'; MFG='制造费用小计'; months=['4月','5月','6月','7月','8月']
def sm(a): return { 'q':sum(x[Q] for x in a),'c':sum(x[C] for x in a),'m':sum(x[MAT] for x in a),'l':sum(x[LAB] for x in a),'f':sum(x[MFG] for x in a)}
S={m:sm(M[m]) for m in months}; allr=[x for m in months for x in M[m]]; miss=sum(x[MAT]==0 for x in allr); small=sum(x[Q]<100 for x in allr); V={m:[x for x in M[m] if x[MAT]>0 and x[Q]>=100] for m in months[-2:]}; VS={m:sm(V[m]) for m in V}
NAVY=RGBColor(20,38,63); INK=RGBColor(35,45,56); MUT=RGBColor(103,116,128); TEAL=RGBColor(19,139,143); CORAL=RGBColor(222,91,75); GOLD=RGBColor(224,164,58); CREAM=RGBColor(247,245,239); WHITE=RGBColor(255,255,255); LINE=RGBColor(220,225,228); GREEN=RGBColor(69,145,95)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
def bg(s,c=CREAM): s.background.fill.solid(); s.background.fill.fore_color.rgb=c
def box(s,x,y,w,h,c,l=None,r=False): sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.color.rgb=l or c; return sh
def tx(s,z,x,y,w,h,sz=14,c=INK,b=False,a=PP_ALIGN.LEFT): q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); q.text_frame.word_wrap=True; p=q.text_frame.paragraphs[0]; p.text=z; p.alignment=a; p.font.name='Microsoft YaHei'; p.font.size=Pt(sz); p.font.bold=b; p.font.color.rgb=c; return q
def title(s,k,h,n=''): tx(s,k,.62,.35,3,.25,10,TEAL,True); tx(s,h,.62,.67,12,.55,27,NAVY,True); tx(s,n,.64,1.2,12,.3,11,MUT); box(s,.62,1.55,12.08,.02,LINE)
def foot(s,z='数据来源：4-8月成本核算汇总表；金额单位：元'): tx(s,z,.62,7.18,12,.18,8.5,MUT)
def kpi(s,x,y,w,l,v,n,c): box(s,x,y,w,1.1,WHITE,LINE,True); box(s,x,y,.08,1.1,c,c); tx(s,l,x+.2,y+.15,w-.3,.22,10,MUT,True); tx(s,v,x+.2,y+.42,w-.3,.38,24,NAVY,True); tx(s,n,x+.2,y+.84,w-.3,.18,9.5,MUT)
def bullet(s,z,x,y,w,c): box(s,x,y+.1,.1,.1,c,c,True); tx(s,z,x+.22,y,w-.22,.32,14,INK)
# 1 cover
s=prs.slides.add_slide(blank); bg(s,NAVY); box(s,0,0,13.33,7.5,NAVY); box(s,0,0,.24,7.5,CORAL); tx(s,'织造车间',.85,1.15,3,.3,15,RGBColor(152,211,209),True); tx(s,'成本管理体系建设',.82,1.62,9,.7,40,WHITE,True); tx(s,'4-8月第一版数据复盘',.86,2.55,8,.4,24,RGBColor(222,229,234)); box(s,.86,3.42,5.2,.04,TEAL); tx(s,'先把数据说清楚，再把成本管起来',.86,3.75,7.5,.4,18,WHITE,True); box(s,8.95,1.22,3.45,4.8,RGBColor(30,57,83),RGBColor(54,82,107),True); tx(s,'本次汇报只回答三个问题',9.35,1.65,2.7,.5,16,WHITE,True); bullet(s,'问题有多大？',9.35,2.45,2.5,CORAL); bullet(s,'哪些数据可以相信？',9.35,3.2,2.8,GOLD); bullet(s,'9-11月先做什么？',9.35,3.95,2.8,TEAL); tx(s,'财务部｜2026年9月',.86,6.78,3,.2,10,RGBColor(177,193,202))
# 2 conclusion
s=prs.slides.add_slide(blank); bg(s); title(s,'01｜先看结论','第一版数据已经足够说明：先治理数据，再谈精细化成本','4-8月共168条记录，最重要的管理动作是把领料和盘点变成可追溯流程。'); kpi(s,.62,1.88,2.8,'样本记录','168 条','4-8月全部记录',TEAL); kpi(s,3.62,1.88,2.8,'原材料缺失',f'{miss} 条',f'占 {miss/len(allr)*100:.1f}%，不能直接做成本判断',CORAL); kpi(s,6.62,1.88,2.8,'小批量记录',f'{small} 条','产量 <100KG',GOLD); kpi(s,9.62,1.88,2.8,'可比样本','7-8月','剔除缺料与小批量',GREEN); box(s,.62,3.35,12.08,2.65,WHITE,LINE,True); tx(s,'管理判断',.92,3.68,1.5,.25,12,TEAL,True); tx(s,'当前不能用全量数据回答“哪个品种最赚钱”。',.92,4.05,10.8,.42,21,NAVY,True); bullet(s,'可以确认：原材料是成本控制的第一抓手，必须绑定订单和盘点。',.94,4.72,10.8,TEAL); bullet(s,'可以推进：9月先把数据完整率和领料流程做实，10月再做工序级分摊。',.94,5.16,10.8,CORAL); foot(s)
# 3 boundary
s=prs.slides.add_slide(blank); bg(s); title(s,'02｜数据边界','先把“能不能比”讲清楚，避免用缺失数据做结论'); box(s,.72,1.92,5,4.55,WHITE,LINE,True); tx(s,'168条记录如何分层',1.02,2.2,3.8,.3,15,NAVY,True); box(s,1.05,2.83,4.3,.66,TEAL,TEAL,True); tx(s,'全部记录 168 条',1.22,3.02,3.9,.25,18,WHITE,True,PP_ALIGN.CENTER); box(s,1.42,3.75,3.56,.66,CORAL,CORAL,True); tx(s,f'原材料金额为0：{miss}条',1.52,3.94,3.35,.25,15,WHITE,True,PP_ALIGN.CENTER); box(s,1.78,4.67,2.84,.66,GOLD,GOLD,True); tx(s,f'产量 <100KG：{small}条',1.86,4.86,2.7,.25,14,WHITE,True,PP_ALIGN.CENTER); box(s,2.1,5.59,2.2,.52,GREEN,GREEN,True); tx(s,'可比样本：7-8月',2.18,5.73,2.05,.2,12,WHITE,True,PP_ALIGN.CENTER); box(s,6.02,1.92,6.68,4.55,RGBColor(240,247,246),RGBColor(197,225,222),True); tx(s,'这意味着什么',6.38,2.25,3,.3,15,NAVY,True); bullet(s,'4-6月没有期初盘点，原材料数据不完整。',6.38,2.95,5.8,CORAL); bullet(s,'全量数据适合找管理问题，不适合直接做品种盈利排名。',6.38,3.55,5.8,GOLD); bullet(s,'7-8月剔除缺料与小批量后，才作为相对可比的趋势参考。',6.38,4.15,5.8,TEAL); bullet(s,'后续每月固定输出：全部样本、可比样本、剔除原因三张清单。',6.38,4.75,5.8,GREEN); foot(s,'口径：缺料=原材料金额为0；小批量=实收数量<100KG；两类记录可重叠。')
# 4 composition
s=prs.slides.add_slide(blank); bg(s); title(s,'03｜成本结构','原材料占比远高于参考区间，成本改善要从“料”开始','行业参考仅作方向性基准，当前数据仍受盘点和领料完整性影响。'); box(s,.75,1.95,7.1,4.6,WHITE,LINE,True); tx(s,'4-8月原材料占比',1.05,2.25,4,.25,14,NAVY,True); pcts=[]
for m in months: pcts.append(S[m]['m']/S[m]['c']*100)
for i,m in enumerate(months): y=2.8+i*.62; pct=pcts[i]; tx(s,m,1.08,y,.55,.22,11,MUT,True); box(s,1.75,y,4.5,.28,RGBColor(233,238,240),RGBColor(233,238,240)); box(s,1.75,y,4.5*pct/100,.28,CORAL,CORAL); tx(s,f'{pct:.1f}%',6.42,y-.01,.8,.22,12,CORAL,True)
box(s,8.18,1.95,4.52,4.6,WHITE,LINE,True); tx(s,'读图结论',8.52,2.27,2.5,.3,15,NAVY,True); tx(s,f'{min(pcts):.0f}%-{max(pcts):.0f}%',8.52,2.85,3.1,.62,34,CORAL,True); tx(s,'原材料成本占比（4-8月）',8.54,3.48,3.6,.3,11,MUT); bullet(s,'原纱采购价、用量和织损是第一优先级。',8.52,4.1,3.6,CORAL); bullet(s,'人工和制造费用占比低，先不要用它们解释全部波动。',8.52,4.78,3.6,TEAL); foot(s)
# 5 comparable
s=prs.slides.add_slide(blank); bg(s); title(s,'04｜可比样本','7-8月有效样本显示：单位成本下降，但不能简单归因于管理改善','有效样本：原材料金额>0且产量≥100KG，用于趋势观察。');
for i,m in enumerate(['7月','8月']): d=VS[m]; x=.9+i*6.0; box(s,x,2.0,5.2,3.8,WHITE,LINE,True); tx(s,m,x+.35,2.3,1.2,.3,19,NAVY,True); tx(s,f'{len(V[m])} 个有效品种',x+.35,2.75,2.5,.25,11,MUT); tx(s,f'{d["c"]/10000:.1f} 万元',x+.35,3.18,3,.4,24,NAVY,True); tx(s,f'{d["q"]/1000:.1f} 吨',x+.35,3.72,2,.25,12,MUT); tx(s,f'{d["c"]/d["q"]:.2f} 元/KG',x+.35,4.22,3,.35,18,TEAL,True); tx(s,f'原材料占比 {d["m"]/d["c"]*100:.1f}%',x+.35,4.8,3,.23,11,MUT); tx(s,f'人工+制造 {((d["l"]+d["f"])/d["c"]*100):.1f}%',x+.35,5.2,3,.23,11,MUT)
box(s,.9,6.25,11.2,.48,RGBColor(255,246,229),RGBColor(242,218,169),True); tx(s,'提示：7月到8月成本下降同时伴随产量和品种结构变化，下一步必须分解价格、用量、产品结构三项影响。',1.15,6.39,10.7,.2,11,INK,True); foot(s)
# 6 action
s=prs.slides.add_slide(blank); bg(s); title(s,'05｜行动计划','9-11月按月只做一件关键事：先完整，再精细，最后预警'); acts=[('9月','把数据补齐','按订单领料；每日核对ERP出库；次品单独登记',CORAL),('10月','把成本算准','工序采集器；按品种分摊人工；标准用量复盘',TEAL),('11月','把异常提前发现','变动率>10%预警；月度成本会；看板试运行',GOLD)]
for i,(mo,hd,bd,c) in enumerate(acts): x=1.0+i*4.05; box(s,x,2.1,3.45,2.45,WHITE,LINE,True); box(s,x,2.1,3.45,.1,c,c); tx(s,mo,x+.3,2.45,1,.3,14,c,True); tx(s,hd,x+.3,2.9,2.6,.3,18,NAVY,True); tx(s,bd,x+.3,3.45,2.7,.75,13,INK)
kpi(s,1.0,5.15,3.45,'9月验收','完整率 ≥90%','按订单领料覆盖主力品种',CORAL); kpi(s,4.95,5.15,3.45,'10月验收','报告 T+2','成本分摊规则可复核',TEAL); kpi(s,8.9,5.15,3.45,'11月验收','异常可预警','变动率>10%自动入清单',GOLD); foot(s)
# 7 asks
s=prs.slides.add_slide(blank); bg(s,NAVY); tx(s,'06｜需要管理层确认',.78,.62,4,.3,12,RGBColor(150,219,216),True); tx(s,'把成本管理从“事后解释”推进到“现场可执行”',.78,1.12,11,.6,29,WHITE,True); box(s,.78,2.1,11.8,.03,RGBColor(74,112,137)); asks=[('01','按订单领料','确认生产部配合，9月先覆盖主力品种。',CORAL),('02','次品独立登记','确认BC级次品单独入库、单独台账。',GOLD),('03','月度成本会','确认每月关账后固定召开，按异常-原因-动作闭环。',TEAL)]
for i,(n,h,b,c) in enumerate(asks): y=2.62+i*1.15; box(s,.82,y,.55,.55,c,c,True); tx(s,n,.82,y+.14,.55,.2,13,WHITE,True,PP_ALIGN.CENTER); tx(s,h,1.65,y+.02,3,.28,18,WHITE,True); tx(s,b,1.65,y+.38,9.9,.3,13,RGBColor(218,229,235))
box(s,.82,6.25,11.65,.58,RGBColor(30,57,83),RGBColor(74,112,137),True); tx(s,'本版数据的价值：把问题边界说清楚，把下一步动作定下来。',1.08,6.43,11.1,.22,15,RGBColor(237,243,245),True,PP_ALIGN.CENTER); tx(s,'财务部｜2026年9月',.82,7.12,3,.18,9,RGBColor(153,176,188))
os.makedirs(os.path.dirname(OUT),exist_ok=True); prs.save(OUT); print(OUT)
