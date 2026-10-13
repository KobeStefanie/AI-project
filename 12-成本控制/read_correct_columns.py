#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
用正确的列位置读取4-8月数据
"""
import openpyxl

def safe_float(val):
    """安全转换为浮点数"""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        if val.startswith('#'):
            return 0.0
        try:
            return float(val.replace(',', ''))
        except:
            return 0.0
    return 0.0

# 定义各月份的列位置
COLUMN_MAP = {
    '4月': {'material': 47, 'labor': 60, 'mfg': 71, 'cost': 73, 'unit': 75},
    '5月': {'material': 47, 'labor': 60, 'mfg': 71, 'cost': 73, 'unit': 75},
    '6月': {'material': 47, 'labor': 60, 'mfg': 71, 'cost': 73, 'unit': 75},
    '7月': {'material': 47, 'labor': 60, 'mfg': 72, 'cost': 74, 'unit': 76},
    '8月': {'material': 47, 'labor': 54, 'mfg': 65, 'cost': 67, 'unit': 69},
}

excel_file = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'
wb = openpyxl.load_workbook(excel_file, data_only=True)

print("=" * 80)
print("4-8月完整数据（使用正确的列位置）")
print("=" * 80)

months_data = {}

for month in ['4月', '5月', '6月', '7月', '8月']:
    ws = wb[month]
    cols = COLUMN_MAP[month]

    varieties = []
    for row_idx in range(3, ws.max_row + 1):
        name = ws.cell(row_idx, 3).value
        if not name or name == '合计':
            continue

        item = {
            '物料名称': name,
            '类型': ws.cell(row_idx, 4).value or '',
            '实收数量': safe_float(ws.cell(row_idx, 9).value),
            '原材料金额': safe_float(ws.cell(row_idx, cols['material']).value),
            '人工小计': safe_float(ws.cell(row_idx, cols['labor']).value),
            '制造费用小计': safe_float(ws.cell(row_idx, cols['mfg']).value),
            '实际成本合计': safe_float(ws.cell(row_idx, cols['cost']).value),
            '实际单位成本': safe_float(ws.cell(row_idx, cols['unit']).value),
        }

        varieties.append(item)

    months_data[month] = varieties

    # 计算汇总
    total_qty = sum(v['实收数量'] for v in varieties)
    total_material = sum(v['原材料金额'] for v in varieties)
    total_labor = sum(v['人工小计'] for v in varieties)
    total_mfg = sum(v['制造费用小计'] for v in varieties)
    total_cost = sum(v['实际成本合计'] for v in varieties)

    print(f"\n{month}:")
    print(f"  品种数: {len(varieties)}")
    print(f"  总产量: {total_qty:,.0f} KG")
    print(f"  总成本: {total_cost/10000:.2f}万元")
    print(f"    原材料: {total_material/10000:.2f}万元 ({total_material/total_cost*100 if total_cost>0 else 0:.1f}%)")
    print(f"    人工: {total_labor/10000:.2f}万元 ({total_labor/total_cost*100 if total_cost>0 else 0:.1f}%)")
    print(f"    制造费用: {total_mfg/10000:.2f}万元 ({total_mfg/total_cost*100 if total_cost>0 else 0:.1f}%)")

print("\n" + "=" * 80)
print("4-8月成本趋势")
print("=" * 80)

for month in ['4月', '5月', '6月', '7月', '8月']:
    varieties = months_data[month]
    total_cost = sum(v['实际成本合计'] for v in varieties)
    total_material = sum(v['原材料金额'] for v in varieties)
    total_labor = sum(v['人工小计'] for v in varieties)
    total_mfg = sum(v['制造费用小计'] for v in varieties)

    print(f"{month}: {len(varieties):2d}品种 | 总成本{total_cost/10000:7.2f}万 = 材料{total_material/10000:7.2f} + 人工{total_labor/10000:5.2f} + 制造{total_mfg/10000:5.2f}")

# 保存完整数据供PPT使用
import pickle
with open(r'd:\AI-项目\12-成本控制\months_data_correct.pkl', 'wb') as f:
    pickle.dump(months_data, f)

print("\n数据已保存到: months_data_correct.pkl")

wb.close()
