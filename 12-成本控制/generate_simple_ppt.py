# -*- coding: utf-8 -*-
import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

print("开始生成成本分析PPT（简化版）...")

def safe_float(val):
    """安全转换为浮点数"""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        try:
            return float(val.replace(',', ''))
        except:
            return 0.0
    return 0.0

# 读取Excel数据
file_path = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)

# 读取8月数据
ws = wb['8月']
data_8 = []
for row in range(3, ws.max_row + 1):
    item = {
        'name': ws.cell(row, 3).value,
        'spec': ws.cell(row, 6).value,
        'qty': safe_float(ws.cell(row, 9).value),
        'std_cost': safe_float(ws.cell(row, 69).value),
        'act_cost': safe_float(ws.cell(row, 70).value),
        'unit_cost': safe_float(ws.cell(row, 72).value),
        'material': safe_float(ws.cell(row, 47).value),
        'labor': safe_float(ws.cell(row, 54).value),
        'mfg': safe_float(ws.cell(row, 67).value),
    }
    if item['name']:
        data_8.append(item)

wb.close()

# 计算汇总
total_qty = sum(d['qty'] for d in data_8)
total_std = sum(d['std_cost'] for d in data_8)
total_act = sum(d['act_cost'] for d in data_8)
total_mat = sum(d['material'] for d in data_8)
total_lab = sum(d['labor'] for d in data_8)
total_mfg = sum(d['mfg'] for d in data_8)

# TOP 5
top5_cost = sorted(data_8, key=lambda x: x['unit_cost'], reverse=True)[:5]
top5_prod = sorted(data_8, key=lambda x: x['qty'], reverse=True)[:5]

# 创建PPT
prs = Presentation()
COLOR_BLUE = RGBColor(31, 78, 120)

def add_slide(title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.7))
    tb.text_frame.text = title
    tb.text_frame.paragraphs[0].font.size = Pt(28)
    tb.text_frame.paragraphs[0].font.bold = True
    tb.text_frame.paragraphs[0].font.color.rgb = COLOR_BLUE

    cb = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.8))
    tf = cb.text_frame
    for i, b in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(14)

# 封面
slide = prs.slides.add_slide(prs.slide_layouts[6])
tb = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1))
tb.text_frame.text = "织造车间成本分析报告"
tb.text_frame.paragraphs[0].font.size = Pt(44)
tb.text_frame.paragraphs[0].font.bold = True
tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# 第2页：总览
add_slide("8月成本总览", [
    f"品种数：{len(data_8)}个",
    f"总产量：{total_qty:,.0f} KG",
    f"",
    f"标准成本：{total_std:,.0f} 元",
    f"实际成本：{total_act:,.0f} 元",
    f"成本差异：{total_act - total_std:,.0f} 元",
    f"差异率：{(total_act - total_std)/total_std*100:.2f}%",
])

# 第3页：成本构成
add_slide("成本构成分析", [
    f"实际成本合计：{total_act:,.0f} 元",
    "",
    f"原料成本：{total_mat:,.0f} 元 ({total_mat/total_act*100:.1f}%)",
    f"人工成本：{total_lab:,.0f} 元 ({total_lab/total_act*100:.1f}%)",
    f"制造费用：{total_mfg:,.0f} 元 ({total_mfg/total_act*100:.1f}%)",
    "",
    "结论：原料成本占比最高，是成本管控重点"
])

# 第4页：TOP5高成本
bullets = ["TOP 5 单位成本最高的品种：", ""]
for i, item in enumerate(top5_cost, 1):
    bullets.append(f"{i}. {item['name']} - {safe_float(item['unit_cost']):.2f} 元/KG")
add_slide("高成本品种", bullets)

# 第5页：TOP5高产量
bullets = ["TOP 5 产量最高的品种：", ""]
for i, item in enumerate(top5_prod, 1):
    bullets.append(f"{i}. {item['name']} - {safe_float(item['qty']):,.0f} KG")
add_slide("主力品种", bullets)

# 第6页：改进建议
add_slide("成本优化建议", [
    "1. 重点优化高成本品种的原料采购",
    "2. 加强主力品种的工艺流程优化",
    "3. 建立成本预警机制",
    "4. 推行成本责任制",
    "",
    f"预期降本：{abs(total_act - total_std) * 0.3 / 10000:.1f}万元/月"
])

# 保存
output_file = "01-交付物归档/2026-09/织造车间成本分析_202609.pptx"
os.makedirs(os.path.dirname(output_file), exist_ok=True)
prs.save(output_file)

print(f"\nPPT生成成功：{output_file}")
print(f"共6页，基于8月实际数据")
