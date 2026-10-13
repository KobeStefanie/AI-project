#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
快速分析Excel核心结构
"""
import openpyxl

# 1. 读取带公式的Excel - 只看sheet结构
print("=" * 80)
print("分析：成本核算表-7月-2版.xlsx")
print("=" * 80)

formula_file = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算表-7月-2版.xlsx'
wb = openpyxl.load_workbook(formula_file, data_only=True, read_only=True)

print(f"\n共有 {len(wb.sheetnames)} 个sheet：")
for i, name in enumerate(wb.sheetnames, 1):
    ws = wb[name]
    print(f"{i:2d}. {name:30s} - {ws.max_row:6d}行")

# 2. 读取成本核算sheet的表头
if '成本核算' in wb.sheetnames:
    print("\n" + "=" * 80)
    print("成本核算sheet - 前20列表头")
    print("=" * 80)
    ws = wb['成本核算']
    for col in range(1, 21):
        h1 = ws.cell(1, col).value or ""
        h2 = ws.cell(2, col).value or ""
        print(f"Col{col:2d}: {h1} / {h2}")

# 3. 读取汇总表8月的完整71列表头
print("\n" + "=" * 80)
print("分析：成本核算汇总-标准成本 - 加工与自产.xlsx (8月)")
print("=" * 80)

summary_file = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'
wb2 = openpyxl.load_workbook(summary_file, data_only=True, read_only=True)

if '8月' in wb2.sheetnames:
    ws = wb2['8月']
    print(f"\n8月表：{ws.max_row}行 x {ws.max_column}列")
    print("\n完整71列表头：")

    output_file = r'd:\AI-项目\12-成本控制\column_headers.txt'
    with open(output_file, 'w', encoding='utf-8') as f:
        for col in range(1, ws.max_column + 1):
            h1 = ws.cell(1, col).value or ""
            h2 = ws.cell(2, col).value or ""
            line = f"Col{col:2d}: [{h1}] / [{h2}]"
            print(line)
            f.write(line + '\n')

    print(f"\n表头已写入: {output_file}")

wb.close()
wb2.close()
print("\n完成！")
