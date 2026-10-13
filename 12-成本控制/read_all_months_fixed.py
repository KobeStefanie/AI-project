#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
完整读取4-8月数据（处理Excel错误值）
"""
import openpyxl

def safe_float(val):
    """安全转换为浮点数，处理Excel错误值"""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        # 处理Excel错误值
        if val.startswith('#'):  # #DIV/0!, #N/A, #VALUE! 等
            return 0.0
        try:
            return float(val.replace(',', ''))
        except:
            return 0.0
    return 0.0

excel_file = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'
wb = openpyxl.load_workbook(excel_file, data_only=True)

print("=" * 80)
print("完整读取4-8月数据")
print("=" * 80)

months_data = {}

for month in ['4月', '5月', '6月', '7月', '8月']:
    if month not in wb.sheetnames:
        continue

    ws = wb[month]

    # 读取所有品种
    varieties = []
    for row_idx in range(3, ws.max_row + 1):
        name = ws.cell(row_idx, 3).value  # Col3: 物料名称
        if not name or name == '合计':
            continue

        # 读取关键字段（使用safe_float处理）
        item = {
            '物料名称': name,
            '类型': ws.cell(row_idx, 4).value or '',
            '实收数量': safe_float(ws.cell(row_idx, 9).value),
            '原材料金额': safe_float(ws.cell(row_idx, 47).value),
            '人工小计': safe_float(ws.cell(row_idx, 54).value),
            '制造费用小计': safe_float(ws.cell(row_idx, 65).value),
            '实际成本合计': safe_float(ws.cell(row_idx, 67).value),
            '实际单位成本': safe_float(ws.cell(row_idx, 69).value),
            '变动率': safe_float(ws.cell(row_idx, 70).value),
        }

        varieties.append(item)

    months_data[month] = varieties

    # 计算月度汇总
    total_qty = sum(v['实收数量'] for v in varieties)
    total_material = sum(v['原材料金额'] for v in varieties)
    total_labor = sum(v['人工小计'] for v in varieties)
    total_mfg = sum(v['制造费用小计'] for v in varieties)
    total_cost = sum(v['实际成本合计'] for v in varieties)

    print(f"\n{month}:")
    print(f"  品种数: {len(varieties)}")
    print(f"  总产量: {total_qty:,.0f} KG")
    print(f"  总成本: {total_cost/10000:.2f}万元")

    if total_cost > 0:
        print(f"    原材料: {total_material/10000:.2f}万元 ({total_material/total_cost*100:.1f}%)")
        print(f"    人工: {total_labor/10000:.2f}万元 ({total_labor/total_cost*100:.1f}%)")
        print(f"    制造费用: {total_mfg/10000:.2f}万元 ({total_mfg/total_cost*100:.1f}%)")
    else:
        print(f"    原材料: {total_material/10000:.2f}万元")
        print(f"    人工: {total_labor/10000:.2f}万元")
        print(f"    制造费用: {total_mfg/10000:.2f}万元")

# 输出详细数据到文件
output_file = r'd:\AI-项目\12-成本控制\months_4_8_complete.txt'
with open(output_file, 'w', encoding='utf-8') as f:
    for month, varieties in months_data.items():
        f.write(f"\n{'='*80}\n")
        f.write(f"{month}完整数据（{len(varieties)}个品种）\n")
        f.write('='*80 + '\n\n')

        # 计算汇总
        total_qty = sum(v['实收数量'] for v in varieties)
        total_material = sum(v['原材料金额'] for v in varieties)
        total_labor = sum(v['人工小计'] for v in varieties)
        total_mfg = sum(v['制造费用小计'] for v in varieties)
        total_cost = sum(v['实际成本合计'] for v in varieties)

        f.write(f"月度汇总:\n")
        f.write(f"  总产量: {total_qty:,.0f} KG\n")
        f.write(f"  总成本: {total_cost:,.0f}元 ({total_cost/10000:.2f}万元)\n")
        f.write(f"  原材料: {total_material:,.0f}元\n")
        f.write(f"  人工: {total_labor:,.0f}元\n")
        f.write(f"  制造费用: {total_mfg:,.0f}元\n\n")

        f.write("品种明细（前10个）:\n")
        for v in varieties[:10]:
            f.write(f"\n{v['物料名称']} ({v['类型']})\n")
            f.write(f"  实收数量: {v['实收数量']:,.0f} KG\n")
            f.write(f"  原材料: {v['原材料金额']:,.0f}元\n")
            f.write(f"  人工: {v['人工小计']:,.0f}元\n")
            f.write(f"  制造费用: {v['制造费用小计']:,.0f}元\n")
            f.write(f"  实际成本: {v['实际成本合计']:,.0f}元\n")
            f.write(f"  单位成本: {v['实际单位成本']:.2f}元/KG\n")
            f.write(f"  变动率: {v['变动率']*100:.1f}%\n")

print(f"\n详细数据已写入: {output_file}")

# 生成4-8月趋势数据
print("\n" + "=" * 80)
print("4-8月成本趋势")
print("=" * 80)

trend_data = []
for month in ['4月', '5月', '6月', '7月', '8月']:
    varieties = months_data[month]
    total_cost = sum(v['实际成本合计'] for v in varieties)
    total_material = sum(v['原材料金额'] for v in varieties)
    total_labor = sum(v['人工小计'] for v in varieties)
    total_mfg = sum(v['制造费用小计'] for v in varieties)

    trend_data.append({
        '月份': month,
        '品种数': len(varieties),
        '总成本': total_cost,
        '原材料': total_material,
        '人工': total_labor,
        '制造费用': total_mfg,
    })

for item in trend_data:
    print(f"{item['月份']}: {item['品种数']:2d}个品种, 总成本{item['总成本']/10000:7.2f}万元 | 原材料{item['原材料']/10000:7.2f} + 人工{item['人工']/10000:5.2f} + 制造{item['制造费用']/10000:5.2f}")

wb.close()
