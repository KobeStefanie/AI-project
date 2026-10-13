#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
读取带公式的Excel，分析数据逻辑
"""
import openpyxl

excel_path = r'C:\Users\Administrator\Desktop\华盛\成本PPT汇报\成本核算汇总-标准成本 - 加工与自产.xlsx'

# 读取带公式的Excel
wb = openpyxl.load_workbook(excel_path, data_only=False)
ws = wb['8月']

output = []
output.append("=" * 80)
output.append("8月sheet数据逻辑分析（读取公式）")
output.append("=" * 80)
output.append("")

# 1. 检查是否有多张表格（通过空行分隔）
output.append("检查表格结构（前50行）：")
output.append("")

for row_idx in range(1, min(51, ws.max_row + 1)):
    # 检查A列是否有值
    a_val = ws.cell(row_idx, 1).value

    # 如果A列是标题或者合计行，标注出来
    if a_val and (isinstance(a_val, str)):
        if any(keyword in str(a_val) for keyword in ['序号', '合计', '备注', '品类', '财务编码']):
            output.append(f"Row{row_idx}: 【{a_val}】<-- 可能是表格分界线")

    # 如果整行为空，也标注
    if all(ws.cell(row_idx, col).value is None for col in range(1, 11)):
        output.append(f"Row{row_idx}: [空行] <-- 可能是表格分隔")

output.append("")
output.append("=" * 80)
output.append("关键列的公式示例（Row3第一条数据）：")
output.append("")

# 2. 检查关键列的公式
key_cols = {
    47: "实际金额小计（原材料）",
    54: "人工小计",
    65: "制造费用小计",
    66: "标准成本合计",
    67: "实际成本合计",
    68: "标准单位成本",
    69: "实际单位成本",
    70: "变动率",
}

for col_idx, col_name in key_cols.items():
    cell = ws.cell(3, col_idx)
    output.append(f"Col{col_idx} ({col_name}):")

    # 如果是公式
    if isinstance(cell.value, str) and cell.value.startswith('='):
        output.append(f"  公式: {cell.value}")
    else:
        output.append(f"  直接值: {cell.value}")
    output.append("")

output.append("=" * 80)
output.append("检查第二张表（原料汇总表）的起始位置：")
output.append("")

# 找到"品类"或"财务编码"关键字，确定第二张表起始行
for row_idx in range(20, min(40, ws.max_row + 1)):
    row_vals = [ws.cell(row_idx, col).value for col in range(1, 8)]
    if any(val and '品类' in str(val) for val in row_vals):
        output.append(f"第二张表起始行: Row{row_idx}")
        output.append(f"表头: {row_vals}")
        output.append("")

        # 读取第二张表的前5行数据
        output.append("第二张表数据样本（前5行）：")
        for i in range(1, 6):
            data_row = row_idx + i
            vals = [ws.cell(data_row, col).value for col in range(1, 8)]
            output.append(f"  Row{data_row}: {vals}")
        break

# 写入文件
output_file = 'd:/AI-项目/12-成本控制/excel_logic_analysis.txt'
with open(output_file, 'w', encoding='utf-8') as f:
    f.write('\n'.join(output))

print(f"Done: {output_file}")
