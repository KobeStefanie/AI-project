import openpyxl,sys
sys.stdout.reconfigure(encoding='utf-8')
p=r'C:\Users\Administrator\Desktop\华盛\成本核算汇总-标准成本 - 加工与自产.xlsx'
w=openpyxl.load_workbook(p,data_only=True,read_only=True)
for sn in w.sheetnames:
 ws=w[sn]; print('===',sn,ws.max_row,ws.max_column,'===')
 for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,8),values_only=True): print(row)
